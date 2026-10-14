"""Renderizado de CV en DOCX y PDF con formato simple compatible con ATS:
una sola columna, sin tablas, sin imágenes ni íconos, títulos de sección estándar y texto seleccionable."""
from __future__ import annotations

import io
from xml.sax.saxutils import escape


def _contact_line(c: dict) -> str:
    return " | ".join(v for v in (c.get("location"), c.get("phone"), c.get("email"), c.get("linkedin"), c.get("website")) if v)


def _period(e: dict) -> str:
    return " - ".join(x for x in (e.get("start"), e.get("end")) if x)


def _edu_line(e: dict) -> tuple[str, str]:
    title = " - ".join(x for x in (e.get("degree"), e.get("field")) if x) or e.get("institution", "")
    meta = ", ".join(x for x in (e.get("institution") if title != e.get("institution") else "", _period(e), e.get("status")) if x)
    return title, meta


def _sections(content: dict):
    """Secuencia de secciones común a DOCX y PDF."""
    if content.get("summary"):
        yield "PERFIL PROFESIONAL", [("p", content["summary"])]
    if content.get("experience"):
        items = []
        for e in content["experience"]:
            items.append(("h", " - ".join(x for x in (e.get("title"), e.get("company")) if x)))
            meta = " | ".join(x for x in (_period(e), e.get("location")) if x)
            if meta:
                items.append(("meta", meta))
            items += [("b", b) for b in e.get("bullets", [])]
        yield "EXPERIENCIA LABORAL", items
    if content.get("education"):
        items = []
        for e in content["education"]:
            t, m = _edu_line(e)
            items.append(("h", t))
            if m:
                items.append(("meta", m))
        yield "FORMACIÓN", items
    if content.get("courses"):
        yield "CURSOS Y CERTIFICACIONES", [
            ("b", ", ".join(x for x in (c.get("name"), c.get("institution") if c.get("institution") not in (c.get("name") or "") else "",
                                        c.get("year") if c.get("year") not in (c.get("name") or "") else "") if x))
            for c in content["courses"]
        ]
    if content.get("skills"):
        yield "HABILIDADES Y HERRAMIENTAS", [("p", ", ".join(content["skills"]))]
    if content.get("soft_skills"):
        yield "COMPETENCIAS", [("p", ", ".join(content["soft_skills"]))]
    if content.get("languages"):
        yield "IDIOMAS", [("p", ", ".join(f"{l.get('language')}{': ' + l['level'] if l.get('level') else ''}" for l in content["languages"]))]


def render_docx(content: dict) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor, Cm

    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Cm(2)
        section.left_margin = section.right_margin = Cm(2.2)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(2)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(content.get("full_name", ""))
    r.bold = True
    r.font.size = Pt(18)
    if content.get("headline"):
        hp = doc.add_paragraph()
        hr = hp.add_run(content["headline"])
        hr.font.size = Pt(11.5)
    contact = _contact_line(content.get("contact", {}))
    if contact:
        doc.add_paragraph(contact)

    for title, items in _sections(content):
        h = doc.add_heading(title, level=2)
        for run in h.runs:
            run.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)
            run.font.size = Pt(12)
            run.font.name = "Calibri"
        h.paragraph_format.space_before = Pt(10)
        for kind, text in items:
            if kind == "h":
                hp = doc.add_paragraph()
                hp.paragraph_format.space_before = Pt(6)
                hp.add_run(text).bold = True
            elif kind == "meta":
                mp = doc.add_paragraph()
                mr = mp.add_run(text)
                mr.italic = True
            elif kind == "b":
                doc.add_paragraph(text, style="List Bullet")
            else:
                doc.add_paragraph(text)
    core = doc.core_properties
    core.title = f"CV {content.get('full_name', '')}"
    core.author = content.get("full_name", "")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_pdf(content: dict) -> bytes:
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2.2 * cm, rightMargin=2.2 * cm, topMargin=2 * cm,
                            bottomMargin=2 * cm, title=f"CV {content.get('full_name', '')}", author=content.get("full_name", ""))
    base = ParagraphStyle("base", fontName="Helvetica", fontSize=10, leading=13.5, alignment=TA_LEFT)
    name = ParagraphStyle("name", parent=base, fontName="Helvetica-Bold", fontSize=17, leading=21)
    headline = ParagraphStyle("headline", parent=base, fontSize=11.5, leading=15)
    sec = ParagraphStyle("sec", parent=base, fontName="Helvetica-Bold", fontSize=11.5, leading=15, spaceBefore=10,
                         spaceAfter=3, textColor=colors.HexColor("#1F2A44"))
    bold = ParagraphStyle("bold", parent=base, fontName="Helvetica-Bold", spaceBefore=5)
    meta = ParagraphStyle("meta", parent=base, fontName="Helvetica-Oblique", textColor=colors.HexColor("#333333"))

    story = [Paragraph(escape(content.get("full_name", "")), name)]
    if content.get("headline"):
        story.append(Paragraph(escape(content["headline"]), headline))
    contact = _contact_line(content.get("contact", {}))
    if contact:
        story.append(Paragraph(escape(contact), base))
    for title, items in _sections(content):
        story.append(Paragraph(escape(title), sec))
        bullets = []

        def flush():
            if bullets:
                story.append(ListFlowable([ListItem(Paragraph(escape(b), base), leftIndent=12) for b in bullets],
                                          bulletType="bullet", start="•", leftIndent=12, bulletFontSize=9))
                bullets.clear()

        for kind, text in items:
            if kind == "b":
                bullets.append(text)
                continue
            flush()
            style = {"h": bold, "meta": meta}.get(kind, base)
            story.append(Paragraph(escape(text), style))
        flush()
    story.append(Spacer(1, 4))
    doc.build(story)
    return buf.getvalue()
