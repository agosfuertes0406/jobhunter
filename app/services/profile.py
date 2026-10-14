"""Carga y análisis del CV: extracción de texto (PDF/DOCX) y construcción del perfil maestro.

Regla central: el perfil solo contiene información presente en el CV o cargada por el usuario.
"""
from __future__ import annotations

import io
import re

from . import text as T
from .ai import AIError, ai_available, complete_json

PROFILE_LIST_FIELDS = ["technical_skills", "tools", "soft_skills", "achievements", "keywords"]


def empty_profile() -> dict:
    return {
        "full_name": "", "headline": "", "email": "", "phone": "", "location": "", "linkedin": "", "website": "",
        "summary": "",
        "experience": [],  # {company,title,location,start,end,responsibilities[],achievements[],tools[]}
        "education": [],  # {institution,degree,field,start,end,status}
        "courses": [],  # {name,institution,year}
        "technical_skills": [], "tools": [], "soft_skills": [],
        "languages": [],  # {language,level}
        "achievements": [], "keywords": [],
    }


def _clean_list(values) -> list[str]:
    out, seen = [], set()
    for v in values or []:
        v = re.sub(r"\s+", " ", str(v)).strip(" •-–·\t")
        if v and T.norm(v) not in seen:
            seen.add(T.norm(v))
            out.append(v)
    return out


def normalize_profile(data: dict | None) -> dict:
    p = empty_profile()
    data = data or {}
    for k in p:
        if k in data and data[k] is not None:
            p[k] = data[k]
    for k in PROFILE_LIST_FIELDS:
        p[k] = _clean_list(p[k])
    exp = []
    for e in p["experience"] or []:
        exp.append({
            "company": str(e.get("company", "")).strip(), "title": str(e.get("title", "")).strip(),
            "location": str(e.get("location", "")).strip(), "start": str(e.get("start", "")).strip(),
            "end": str(e.get("end", "")).strip(), "responsibilities": _clean_list(e.get("responsibilities")),
            "achievements": _clean_list(e.get("achievements")), "tools": _clean_list(e.get("tools")),
        })
    p["experience"] = exp
    p["education"] = [
        {k: str(e.get(k, "")).strip() for k in ("institution", "degree", "field", "start", "end", "status")}
        for e in p["education"] or []
    ]
    p["courses"] = [{k: str(c.get(k, "")).strip() for k in ("name", "institution", "year")} for c in p["courses"] or []]
    p["languages"] = [{k: str(l.get(k, "")).strip() for k in ("language", "level")} for l in p["languages"] or []]
    return p


def profile_years(profile: dict) -> float:
    return T.total_experience_years(profile.get("experience", []))


def profile_corpus(profile: dict) -> str:
    """Todo el texto del perfil, para comparar contra requisitos de avisos."""
    parts = [profile.get("headline", ""), profile.get("summary", "")]
    for e in profile.get("experience", []):
        parts += [e.get("title", ""), e.get("company", "")] + e.get("responsibilities", []) + e.get("achievements", []) + e.get("tools", [])
    for ed in profile.get("education", []):
        parts += [ed.get("degree", ""), ed.get("field", ""), ed.get("institution", ""), ed.get("status", "")]
    for c in profile.get("courses", []):
        parts += [c.get("name", ""), c.get("institution", "")]
    for k in PROFILE_LIST_FIELDS:
        parts += profile.get(k, [])
    parts += [f"{l.get('language','')} {l.get('level','')}" for l in profile.get("languages", [])]
    return "\n".join(p for p in parts if p)


def profile_skills(profile: dict) -> list[str]:
    skills = list(profile.get("technical_skills", [])) + list(profile.get("tools", []))
    for e in profile.get("experience", []):
        skills += e.get("tools", [])
    return _clean_list(skills)


# ---------------- extracción de texto ----------------
class CVReadError(Exception):
    pass


def extract_text(filename: str, content: bytes) -> tuple[str, str]:
    name = filename.lower()
    if name.endswith(".pdf") or content[:4] == b"%PDF":
        import pdfplumber

        try:
            with pdfplumber.open(io.BytesIO(content)) as pdf:
                pages = [(page.extract_text(x_tolerance=1.5) or "") for page in pdf.pages]
        except Exception as exc:  # pdfminer lanza varios tipos
            raise CVReadError(f"No se pudo leer el PDF: {exc}") from exc
        text = "\n".join(pages).strip()
        if len(text) < 40:
            raise CVReadError("El PDF no contiene texto seleccionable (¿es una imagen escaneada?). Subí un PDF con texto o un DOCX.")
        return text, "application/pdf"
    if name.endswith(".docx"):
        import docx

        try:
            d = docx.Document(io.BytesIO(content))
        except Exception as exc:
            raise CVReadError(f"No se pudo leer el DOCX: {exc}") from exc
        lines = []
        for para in d.paragraphs:
            t = para.text.strip()
            if t:
                style = (para.style.name or "").lower() if para.style is not None else ""
                lines.append(("• " + t) if "list" in style else t)
        for table in d.tables:
            for row in table.rows:
                seen = []
                for cell in row.cells:
                    ct = cell.text.strip()
                    if ct and ct not in seen:
                        seen.append(ct)
                lines.extend(seen)
        text = "\n".join(lines).strip()
        if len(text) < 40:
            raise CVReadError("El DOCX parece vacío.")
        return text, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    raise CVReadError("Formato no soportado. Subí un archivo PDF o DOCX.")


# ---------------- análisis con IA ----------------
_STR = {"type": "string"}
_STRS = {"type": "array", "items": {"type": "string"}}


def _obj(props: dict) -> dict:
    return {"type": "object", "properties": props, "required": list(props), "additionalProperties": False}


PROFILE_SCHEMA = _obj({
    "full_name": _STR, "headline": _STR, "email": _STR, "phone": _STR, "location": _STR, "linkedin": _STR,
    "website": _STR, "summary": _STR,
    "experience": {"type": "array", "items": _obj({
        "company": _STR, "title": _STR, "location": _STR, "start": _STR, "end": _STR,
        "responsibilities": _STRS, "achievements": _STRS, "tools": _STRS,
    })},
    "education": {"type": "array", "items": _obj({
        "institution": _STR, "degree": _STR, "field": _STR, "start": _STR, "end": _STR, "status": _STR,
    })},
    "courses": {"type": "array", "items": _obj({"name": _STR, "institution": _STR, "year": _STR})},
    "technical_skills": _STRS, "tools": _STRS, "soft_skills": _STRS,
    "languages": {"type": "array", "items": _obj({"language": _STR, "level": _STR})},
    "achievements": _STRS, "keywords": _STRS,
})

PROFILE_SYSTEM = """Sos un extractor de datos de currículums. Convertís el texto de un CV en un perfil estructurado.
Reglas estrictas:
- Copiá solo información que esté explícitamente en el texto. No infieras, no completes, no inventes empresas,
  fechas, títulos, herramientas, métricas ni niveles de idioma. Si un dato no está, dejá "" o una lista vacía.
- Conservá nombres de empresas, cargos, instituciones y fechas tal como aparecen (fechas como "03/2021", "Mar 2021", "Actualidad").
- "responsibilities": funciones descritas en cada experiencia, en frases breves fieles al original.
- "achievements": solo logros con resultado verificable en el texto (preferentemente cuantificados), copiando las cifras exactas.
- "technical_skills": conocimientos técnicos/disciplinares mencionados; "tools": software y herramientas mencionadas;
  "soft_skills": habilidades blandas declaradas en el CV.
- "keywords": 10 a 25 términos profesionales relevantes que aparecen literalmente en el CV.
- "status" en educación: "Completo", "En curso", "Incompleto" solo si el CV lo indica; si no, "".
"""


def extract_profile_ai(cv_text: str) -> dict:
    data = complete_json(PROFILE_SYSTEM, f"Texto del CV:\n<cv>\n{cv_text}\n</cv>", PROFILE_SCHEMA, effort="medium")
    return normalize_profile(data)


# ---------------- análisis heurístico (sin IA) ----------------
SECTION_PATTERNS = {
    "summary": r"(perfil( profesional)?|resumen( profesional)?|sobre mi|acerca de mi|objetivo( profesional)?|summary|profile|about me)",
    "experience": r"(experiencia( laboral| profesional)?|antecedentes laborales|historial laboral|trayectoria( laboral| profesional)?|work experience|experience|employment)",
    "education": r"(formacion( academica)?|educacion|estudios|formacion profesional|education|academic background)",
    "courses": r"(cursos|capacitaciones|certificaciones|cursos y certificaciones|formacion complementaria|courses|certifications|training)",
    "skills": r"(habilidades( tecnicas)?|conocimientos( tecnicos| informaticos)?|herramientas|competencias( tecnicas)?|aptitudes|skills|technical skills|tools)",
    "soft": r"(habilidades blandas|competencias blandas|soft skills)",
    "languages": r"(idiomas|languages)",
    "achievements": r"(logros|achievements|logros destacados)",
}

DATE_RANGE = re.compile(
    r"((?:[A-Za-zñÑáéíóú]{3,10}\.?\s*(?:de\s+)?)?\d{4}|\d{1,2}[/\-.]\d{4})\s*(?:-|–|—|a|al|hasta|to)\s*"
    r"((?:[A-Za-zñÑáéíóú]{3,10}\.?\s*(?:de\s+)?)?\d{4}|\d{1,2}[/\-.]\d{4}|actualidad|presente|actual|hoy|present|current)",
    re.IGNORECASE,
)


def _section_of(line: str) -> str | None:
    n = T.norm(line).rstrip(":").strip()
    if not n or len(n) > 45:
        return None
    for key in ("soft", "courses", "achievements", "languages", "education", "experience", "skills", "summary"):
        if re.fullmatch(SECTION_PATTERNS[key], n):
            return key
    return None


def _split_items(line: str) -> list[str]:
    return [x.strip() for x in re.split(r"[,;|•·/]| - ", line) if x.strip()]


def extract_profile_heuristic(cv_text: str) -> tuple[dict, list[str]]:
    warnings = ["Análisis local sin IA: revisá y corregí cada sección del perfil antes de usarlo."]
    p = empty_profile()
    lines = [l.strip() for l in cv_text.splitlines()]
    lines = [l for l in lines if l]

    m = re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", cv_text)
    p["email"] = m.group(0) if m else ""
    m = re.search(r"(\+?\d[\d\s().-]{7,}\d)", cv_text)
    p["phone"] = m.group(1).strip() if m else ""
    m = re.search(r"(https?://)?(www\.)?linkedin\.com/in/[\w\-%]+/?", cv_text, re.I)
    p["linkedin"] = m.group(0) if m else ""
    for l in lines[:5]:
        if not re.search(r"[@\d]|linkedin|http", l, re.I) and 2 <= len(l.split()) <= 6 and not _section_of(l):
            p["full_name"] = l.title() if l.isupper() else l
            break
    if p["full_name"]:
        idx = lines.index(next(l for l in lines if l.lower() == p["full_name"].lower() or l == p["full_name"]))
        if idx + 1 < len(lines) and not _section_of(lines[idx + 1]) and not re.search(r"[@]|\d{6,}", lines[idx + 1]):
            p["headline"] = lines[idx + 1][:120]
    m = re.search(r"(CABA|Buenos Aires|Capital Federal|C\.A\.B\.A\.?)[^\n|]{0,40}", cv_text)
    p["location"] = m.group(0).strip() if m else ""

    sections: dict[str, list[str]] = {}
    current = None
    for l in lines:
        sec = _section_of(l)
        if sec:
            current = sec
            sections.setdefault(sec, [])
            continue
        if current:
            sections[current].append(l)

    p["summary"] = " ".join(sections.get("summary", []))[:1200]

    # Experiencia: cada rango de fechas abre una entrada.
    exp_lines = sections.get("experience", [])
    entries, cur = [], None
    for i, l in enumerate(exp_lines):
        m = DATE_RANGE.search(l)
        if m:
            header = DATE_RANGE.sub("", l).strip(" |-–—,()")
            if not header and cur is None and i > 0:
                header = exp_lines[i - 1]
            cur = {"company": "", "title": "", "location": "", "start": m.group(1).strip(), "end": m.group(2).strip(),
                   "responsibilities": [], "achievements": [], "tools": [], "_header": header}
            entries.append(cur)
        elif cur is not None:
            is_bullet = l[:1] in "•-–*·▪●" or (len(l) > 60)
            if is_bullet:
                item = l.lstrip("•-–*·▪● ").strip()
                cur["responsibilities"].append(item)
                if re.search(r"\d+\s?%|\$\s?\d|\d+\s?(clientes|personas|empleados|colaboradores|procesos|proyectos|ingresos|busquedas|búsquedas)", item, re.I):
                    cur["achievements"].append(item)
            elif not cur["_header"]:
                cur["_header"] = l
            elif not cur["title"] or not cur["company"]:
                cur["_header"] += " | " + l
    for e in entries:
        parts = [x.strip() for x in re.split(r"\s[|–—-]\s|\s+en\s+|,\s", e.pop("_header")) if x.strip()]
        if parts:
            e["title"] = parts[0]
        if len(parts) > 1:
            e["company"] = parts[1]
        if len(parts) > 2:
            e["location"] = parts[2]
        e["tools"] = T.find_vocab(" ".join(e["responsibilities"]), T.SKILL_VOCAB)
    p["experience"] = entries
    if exp_lines and not entries:
        warnings.append("No se pudieron identificar períodos en la experiencia laboral: completalos manualmente.")

    # Educación
    edu_lines = sections.get("education", [])
    edu, cur = [], None
    for l in edu_lines:
        m = DATE_RANGE.search(l) or re.search(r"\b(19|20)\d{2}\b", l)
        clean = DATE_RANGE.sub("", l).strip(" |-–—,()") if DATE_RANGE.search(l) else l
        if cur is None or (m and cur.get("_has_date")):
            cur = {"institution": "", "degree": clean, "field": "", "start": "", "end": "", "status": "", "_has_date": bool(m)}
            edu.append(cur)
        elif not cur["institution"]:
            cur["institution"] = clean
        if m:
            cur["_has_date"] = True
            if isinstance(m, re.Match) and m.re is DATE_RANGE:
                cur["start"], cur["end"] = m.group(1), m.group(2)
            else:
                cur["end"] = m.group(0)
        if re.search(r"en curso|cursando|actualidad", l, re.I):
            cur["status"] = "En curso"
    for e in edu:
        e.pop("_has_date", None)
    p["education"] = edu

    p["courses"] = [{"name": l.lstrip("•-–* "), "institution": "", "year": (re.search(r"(19|20)\d{2}", l) or [""])[0]}
                    for l in sections.get("courses", [])]

    skill_items = []
    for l in sections.get("skills", []):
        skill_items += _split_items(l.lstrip("•-–* "))
    vocab_hits = set(T.find_vocab(cv_text, T.SKILL_VOCAB))
    p["tools"] = [s for s in skill_items if T.norm(s) in vocab_hits or len(s.split()) <= 3]
    p["technical_skills"] = [s for s in skill_items if s not in p["tools"]]
    folded = T.strip_accents(cv_text).lower()
    for v in sorted(vocab_hits):
        if not any(T.norm(v) == T.norm(s) for s in skill_items):
            # Detectado literalmente en el texto del CV: se agrega con la grafía original (no es inventado).
            m = re.search(r"(?<![a-z0-9])" + re.escape(v) + r"(?![a-z0-9])", folded)
            original = cv_text[m.start():m.end()] if m and len(folded) == len(cv_text) else v
            p["technical_skills"].append(original[:1].upper() + original[1:] if original.islower() else original)
    soft = []
    for l in sections.get("soft", []):
        soft += _split_items(l.lstrip("•-–* "))
    soft += [s for s in T.find_vocab(cv_text, T.SOFT_SKILLS_VOCAB) if not any(T.norm(s) == T.norm(x) for x in soft)]
    p["soft_skills"] = soft

    langs = []
    for l in sections.get("languages", []) or [l for l in lines if re.search(r"ingl[eé]s|english|portugu[eé]s", l, re.I)]:
        for item in re.split(r"[,;|•]", l):
            m = re.match(r"\s*([A-Za-zÁÉÍÓÚáéíóúñ]+)\s*[:\-–]?\s*(.*?)\s*$", item)
            if m and T.norm(m.group(1)) in T.LANG_NAMES:
                level = m.group(2).strip()
                if level.startswith("(") and level.endswith(")"):
                    level = level[1:-1]
                langs.append({"language": T.LANG_NAMES[T.norm(m.group(1))], "level": level})
    p["languages"] = langs

    ach = list(sections.get("achievements", []))
    for e in entries:
        ach += e["achievements"]
    p["achievements"] = [a.lstrip("•-–* ") for a in ach]

    freq: dict[str, int] = {}
    for t in T.tokens(" ".join(sections.get("experience", []) + sections.get("summary", []))):
        if len(t) > 3:
            freq[t] = freq.get(t, 0) + 1
    p["keywords"] = list(dict.fromkeys(
        list(vocab_hits) + [w for w, c in sorted(freq.items(), key=lambda x: -x[1]) if c >= 2][:15]
    ))[:25]
    if not entries:
        warnings.append("No se detectó la sección de experiencia. Verificá los títulos de sección del CV.")
    return normalize_profile(p), warnings


def analyze_cv(cv_text: str) -> tuple[dict, str, list[str]]:
    """Devuelve (perfil, método, advertencias). Usa Claude si está configurado; si falla, heurística local."""
    if ai_available():
        try:
            return extract_profile_ai(cv_text), "claude", []
        except AIError as exc:
            profile, warnings = extract_profile_heuristic(cv_text)
            return profile, "heuristico", [f"La IA no estuvo disponible ({exc}); se usó el análisis local."] + warnings
    profile, warnings = extract_profile_heuristic(cv_text)
    return profile, "heuristico", warnings
