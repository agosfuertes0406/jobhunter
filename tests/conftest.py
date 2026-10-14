"""Fixtures de prueba. TODOS los datos de este directorio son FICTICIOS (datos de prueba), creados
solo para verificar el funcionamiento; no representan personas, empresas ni avisos reales."""
import io
import os
import sys
import tempfile
from pathlib import Path

import pytest

_tmp = tempfile.mkdtemp(prefix="jobhunter-test-")
os.environ["DATA_DIR"] = _tmp
os.environ["DATABASE_URL"] = f"sqlite:///{_tmp}/test.db"
os.environ["AI_ENABLED"] = "false"
os.environ["JOBHUNTER_ENV_FILE"] = "/nonexistent"
for k in ("SMTP_HOST", "IMAP_HOST", "JOOBLE_API_KEY", "CAREERJET_API_KEY"):
    os.environ.pop(k, None)
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

FIXTURES = Path(__file__).parent / "fixtures"

CV_LINES = [
    ("title", "Lucía Prueba Ficticia"),
    ("p", "Analista de Recursos Humanos"),
    ("p", "lucia.prueba@example.com | +54 11 5555-0000 | CABA, Buenos Aires"),
    ("h", "PERFIL PROFESIONAL"),
    ("p", "Analista de RR.HH. con experiencia en reclutamiento y selección, onboarding y liquidación de sueldos."),
    ("h", "EXPERIENCIA LABORAL"),
    ("p", "Analista de Selección | Empresa Ficticia Uno SA | 03/2021 - Actualidad"),
    ("b", "Gestión integral de búsquedas de perfiles administrativos y técnicos con entrevistas por competencias."),
    ("b", "Diseño del proceso de onboarding para nuevos ingresos."),
    ("b", "Reduje el tiempo promedio de cobertura de búsquedas un 25% en 2023."),
    ("b", "Carga y seguimiento de candidatos en el ATS y reportes en Excel avanzado."),
    ("p", "Asistente de Administración de Personal | Comercial Ficticia SRL | 02/2019 - 02/2021"),
    ("b", "Liquidación de sueldos de 120 empleados con Tango Gestión."),
    ("b", "Control de novedades, licencias y legajos."),
    ("h", "FORMACIÓN"),
    ("p", "Licenciatura en Relaciones del Trabajo - Universidad de Prueba | 2014 - 2019"),
    ("h", "CURSOS"),
    ("b", "People Analytics - Instituto Ficticio 2022"),
    ("h", "HABILIDADES"),
    ("p", "Excel avanzado, Power BI, Tango Gestión, LinkedIn Recruiter, Google Workspace"),
    ("h", "IDIOMAS"),
    ("p", "Inglés: Intermedio (B2)"),
]


def make_docx() -> bytes:
    import docx

    d = docx.Document()
    for kind, text in CV_LINES:
        if kind == "title":
            d.add_heading(text, 0)
        elif kind == "h":
            d.add_heading(text, 1)
        elif kind == "b":
            d.add_paragraph(text, style="List Bullet")
        else:
            d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_pdf() -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    y = 800
    for kind, text in CV_LINES:
        c.setFont("Helvetica-Bold" if kind in ("title", "h") else "Helvetica", 10)
        c.drawString(40, y, ("• " if kind == "b" else "") + text)
        y -= 16
    c.save()
    return buf.getvalue()


@pytest.fixture
def cv_docx():
    return make_docx()


@pytest.fixture
def cv_pdf():
    return make_pdf()


@pytest.fixture
def client():
    from fastapi.testclient import TestClient

    from app.main import app

    with TestClient(app) as c:
        c.headers.update({"X-Requested-With": "JobHunter"})
        yield c
